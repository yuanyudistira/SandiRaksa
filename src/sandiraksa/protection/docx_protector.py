"""
DOCX file protection service.

Detects and replaces PII entities in Word documents with tokens.
Supports paragraphs, tables, headers, footers.
"""

from __future__ import annotations

import gc
import logging
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Callable

from sandiraksa.protection.tokenizer import Tokenizer, TokenizerFactory

logger = logging.getLogger(__name__)


# Matches dates like "14-Feb-1988", "14/02/1988", "1988-02-14"
_DATE_HINT = re.compile(
    r"\d{1,2}[\s\-/][A-Za-z0-9]{1,9}[\s\-/]\d{2,4}"
    r"|\d{4}[\-/]\d{1,2}[\-/]\d{1,2}"
)


def _looks_like_date(text: str) -> bool:
    """Quick check that a value resembles a date."""
    return bool(_DATE_HINT.search(text.strip()))


@dataclass
class DetectedEntity:
    """A detected PII entity in document."""
    
    entity_type: str
    value: str
    location: str  # e.g., "paragraph", "table", "header", "footer"
    context: str = ""  # Surrounding text for preview
    confidence: float = 0.8


@dataclass
class DocxProtectionResult:
    """Result of protecting a DOCX file."""
    
    success: bool
    output_path: str | None = None
    error_message: str | None = None
    entities_protected: int = 0
    paragraphs_processed: int = 0
    tables_processed: int = 0
    duration_seconds: float = 0.0


@dataclass
class DocxProtectionProgress:
    """Progress information during protection."""
    
    phase: str = "Initializing"
    entities_found: int = 0
    entities_protected: int = 0
    current_section: str = ""


class DocxProtector:
    """
    Protects DOCX files by detecting and replacing PII with tokens.
    
    Features:
    - Scans paragraphs, tables, headers, footers
    - Auto-detects PII entities using regex patterns
    - Preserves document formatting
    - Uses consistent tokenization (same value = same token)
    """
    
    # Regex patterns for common PII types
    PATTERNS = {
        "EMAIL": r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b',
        "PHONE_NUMBER": r'\b(?:\+62|62|0)[\s.-]?(?:\d{2,4})[\s.-]?(?:\d{3,4})[\s.-]?(?:\d{3,4})\b',
        "ID_NIK": r'\b[1-9]\d{15}\b',
        "ID_NPWP": r'\b\d{2}\.?\d{3}\.?\d{3}\.?\d[-.]?\d{3}\.?\d{3}\b',
        # Credit card: 4 groups of 4 digits with a CONSISTENT separator
        # (all spaces, all dashes, or none). Uses a backreference so mixed
        # patterns like "3174-19880214-1001" (a NIK-style ID) do NOT match.
        "CREDIT_CARD": r'\b\d{4}([\s-]?)\d{4}\1\d{4}\1\d{4}\b',
        "IP_ADDRESS": r'\b(?:\d{1,3}\.){3}\d{1,3}\b',
        "DATE": r'\b(?:\d{1,2}[-/]\d{1,2}[-/]\d{2,4}|\d{4}[-/]\d{1,2}[-/]\d{1,2})\b',
        "URL": r'https?://[^\s<>"{}|\\^`\[\]]+',
    }
    
    # Context patterns for names
    CONTEXT_PATTERNS = {
        "PERSON": r'(?:(?:Nama|Name|Pasien|Patient|Dr\.?|Dokter|Bpk\.?|Ibu|Sdr\.?|Sdri\.?)\s*[:\s]\s*)([A-Z][a-z]+(?:\s+[A-Z][a-z]+){1,3})',
    }
    
    def __init__(
        self,
        project_id: str,
        entity_types: list[str] | None = None,
    ) -> None:
        """
        Initialize protector.
        
        Args:
            project_id: Project ID for tokenization context.
            entity_types: List of entity types to detect. If None, detect all.
        """
        self._project_id = project_id
        self._entity_types = entity_types
        self._tokenizer = TokenizerFactory.get_tokenizer(project_id)
    
    def detect_entities(self, file_path: Path | str) -> list[DetectedEntity]:
        """
        Detect PII entities in a DOCX file.
        
        Args:
            file_path: Path to DOCX file.
            
        Returns:
            List of detected entities.
        """
        import docx
        
        gc_was_enabled = gc.isenabled()
        gc.disable()
        
        try:
            file_path = Path(file_path)
            doc = docx.Document(file_path)
            entities = []
            
            # Scan paragraphs
            for para in doc.paragraphs:
                text = para.text
                if text.strip():
                    para_entities = self._detect_in_text(text, "paragraph")
                    entities.extend(para_entities)
            
            # Scan tables with header context. The first row is treated as the
            # header row; each cell value is classified using its column header
            # (e.g. header "Nama Mock" -> the value is a PERSON; header
            # "Tanggal Lahir" -> the value is a DATE_OF_BIRTH).
            for table in doc.tables:
                rows = list(table.rows)
                if not rows:
                    continue

                headers = [c.text.strip() for c in rows[0].cells]

                for row in rows[1:]:
                    for col_idx, cell in enumerate(row.cells):
                        text = cell.text
                        if not text.strip():
                            continue
                        header = headers[col_idx] if col_idx < len(headers) else ""
                        cell_entities = self._detect_in_cell(text, header)
                        entities.extend(cell_entities)
            
            # Scan headers
            for section in doc.sections:
                header = section.header
                for para in header.paragraphs:
                    if para.text.strip():
                        header_entities = self._detect_in_text(para.text, "header")
                        entities.extend(header_entities)
                
                # Scan footer
                footer = section.footer
                for para in footer.paragraphs:
                    if para.text.strip():
                        footer_entities = self._detect_in_text(para.text, "footer")
                        entities.extend(footer_entities)
            
            # Deduplicate by value
            seen = set()
            unique_entities = []
            for e in entities:
                if e.value not in seen:
                    seen.add(e.value)
                    unique_entities.append(e)
            
            return unique_entities
            
        except Exception as e:
            logger.error(f"Error detecting entities in DOCX: {e}")
            return []
        finally:
            if gc_was_enabled:
                gc.enable()

    def _detect_in_cell(self, text: str, header: str) -> list[DetectedEntity]:
        """
        Detect entities in a table cell using the column header as context.

        If the header maps to a sensitive entity type (e.g. "Nama Mock" ->
        PERSON, "Tanggal Lahir" -> DATE_OF_BIRTH), the whole cell value is
        treated as that entity. Otherwise, fall back to regex-based detection.
        """
        text = text.strip()
        if not text:
            return []

        from sandiraksa.detection.recognizers.context_classifier import (
            classify_label,
        )

        entity_type = classify_label(header) if header else None

        if entity_type:
            # Respect entity_types filter if configured
            if self._entity_types and entity_type not in self._entity_types:
                return []

            # For date-of-birth columns, ensure the value looks like a date
            if entity_type == "DATE_OF_BIRTH" and not _looks_like_date(text):
                return []

            # For NIK/NPWP etc. that already have their own regex, still let the
            # regex detectors handle the exact value below; but for PERSON and
            # DATE_OF_BIRTH the whole cell is the value.
            if entity_type in ("PERSON", "DATE_OF_BIRTH", "LOCATION"):
                return [DetectedEntity(
                    entity_type=entity_type,
                    value=text,
                    location="table",
                    context=f"{header}: {text}",
                    confidence=0.9,
                )]

        # Fall back to regex detection for the cell content
        return self._detect_in_text(text, "table")
    
    def _detect_in_text(self, text: str, location: str) -> list[DetectedEntity]:
        """Detect entities in a text string."""
        entities = []
        
        # Apply standard patterns
        for entity_type, pattern in self.PATTERNS.items():
            if self._entity_types and entity_type not in self._entity_types:
                continue
            
            for match in re.finditer(pattern, text, re.IGNORECASE):
                value = match.group(0)
                # Get context (surrounding 30 chars)
                start = max(0, match.start() - 30)
                end = min(len(text), match.end() + 30)
                context = text[start:end]
                
                entities.append(DetectedEntity(
                    entity_type=entity_type,
                    value=value,
                    location=location,
                    context=f"...{context}...",
                ))
        
        # Apply context patterns (for names)
        for entity_type, pattern in self.CONTEXT_PATTERNS.items():
            if self._entity_types and entity_type not in self._entity_types:
                continue
            
            for match in re.finditer(pattern, text, re.IGNORECASE):
                value = match.group(1)  # Capture group
                start = max(0, match.start() - 20)
                end = min(len(text), match.end() + 20)
                context = text[start:end]
                
                entities.append(DetectedEntity(
                    entity_type=entity_type,
                    value=value,
                    location=location,
                    context=f"...{context}...",
                ))

        # Context-aware recognizers for names and birth dates in narrative text
        entities.extend(self._detect_context_aware(text, location))

        # Global user-defined custom patterns (e.g. hospital MR numbers)
        entities.extend(self._detect_custom_patterns(text, location))

        return entities

    def _detect_custom_patterns(self, text: str, location: str) -> list[DetectedEntity]:
        """Detect matches from global user-defined custom patterns."""
        found: list[DetectedEntity] = []
        try:
            from sandiraksa.detection.custom_patterns import find_custom_matches

            for m in find_custom_matches(text):
                found.append(DetectedEntity(
                    entity_type=m.label,
                    value=m.value,
                    location=location,
                    context=text[:60],
                    confidence=0.95,
                ))
        except Exception as e:
            logger.warning(f"Custom pattern detection failed: {e}")
        return found

    def _detect_context_aware(self, text: str, location: str) -> list[DetectedEntity]:
        """Use context-aware recognizers for PERSON and DATE_OF_BIRTH."""
        found: list[DetectedEntity] = []
        try:
            from sandiraksa.detection.recognizers.id_person import (
                IndonesianPersonRecognizer,
            )
            from sandiraksa.detection.recognizers.id_dob import (
                DateOfBirthRecognizer,
            )

            recognizers = []
            if not self._entity_types or "PERSON" in self._entity_types:
                recognizers.append(IndonesianPersonRecognizer())
            if not self._entity_types or "DATE_OF_BIRTH" in self._entity_types:
                recognizers.append(DateOfBirthRecognizer())

            for rec in recognizers:
                for r in rec.analyze(text, rec.supported_entities):
                    found.append(DetectedEntity(
                        entity_type=r.entity_type,
                        value=r.text,
                        location=location,
                        context=text[:60],
                        confidence=r.score,
                    ))
        except Exception as e:
            logger.warning(f"Context-aware detection failed: {e}")
        return found
    
    def protect_file(
        self,
        input_path: Path | str,
        output_path: Path | str,
        entities_to_protect: list[DetectedEntity] | None = None,
        progress_callback: Callable[[DocxProtectionProgress], None] | None = None,
    ) -> DocxProtectionResult:
        """
        Protect a DOCX file by replacing detected entities with tokens.
        
        Args:
            input_path: Path to input DOCX file.
            output_path: Path for protected output file.
            entities_to_protect: Specific entities to protect. If None, detect all.
            progress_callback: Optional callback for progress updates.
            
        Returns:
            DocxProtectionResult with status and statistics.
        """
        import docx
        
        gc_was_enabled = gc.isenabled()
        gc.disable()
        
        input_path = Path(input_path)
        output_path = Path(output_path)
        start_time = datetime.now()
        progress = DocxProtectionProgress()
        
        entities_protected = 0
        paragraphs_processed = 0
        tables_processed = 0
        
        try:
            # Detect entities if not provided
            if entities_to_protect is None:
                entities_to_protect = self.detect_entities(input_path)
            
            if not entities_to_protect:
                return DocxProtectionResult(
                    success=True,
                    output_path=str(output_path),
                    entities_protected=0,
                    error_message="Tidak ada data sensitif ditemukan",
                )
            
            # Build value -> token map
            value_to_token = {}
            for entity in entities_to_protect:
                if entity.value not in value_to_token:
                    token = self._tokenizer.get_or_create_token(entity.entity_type, entity.value)
                    value_to_token[entity.value] = token
            
            # Load document
            progress.phase = "Loading document"
            if progress_callback:
                progress_callback(progress)
            
            doc = docx.Document(input_path)
            
            # Process paragraphs
            progress.phase = "Processing paragraphs"
            for para in doc.paragraphs:
                if para.text.strip():
                    count = self._replace_in_paragraph(para, value_to_token)
                    entities_protected += count
                    paragraphs_processed += 1
                    
                    progress.entities_protected = entities_protected
                    progress.current_section = f"Paragraph {paragraphs_processed}"
                    if progress_callback:
                        progress_callback(progress)
            
            # Process tables
            progress.phase = "Processing tables"
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            if para.text.strip():
                                count = self._replace_in_paragraph(para, value_to_token)
                                entities_protected += count
                tables_processed += 1
                
                progress.current_section = f"Table {tables_processed}"
                if progress_callback:
                    progress_callback(progress)
            
            # Process headers/footers
            progress.phase = "Processing headers/footers"
            for section in doc.sections:
                # Header
                for para in section.header.paragraphs:
                    if para.text.strip():
                        count = self._replace_in_paragraph(para, value_to_token)
                        entities_protected += count
                
                # Footer
                for para in section.footer.paragraphs:
                    if para.text.strip():
                        count = self._replace_in_paragraph(para, value_to_token)
                        entities_protected += count
            
            # Save document
            progress.phase = "Saving document"
            if progress_callback:
                progress_callback(progress)
            
            output_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(output_path)
            
            duration = (datetime.now() - start_time).total_seconds()
            
            logger.info(
                f"DOCX protection complete: {entities_protected} entities protected "
                f"in {duration:.2f}s"
            )
            
            return DocxProtectionResult(
                success=True,
                output_path=str(output_path),
                entities_protected=entities_protected,
                paragraphs_processed=paragraphs_processed,
                tables_processed=tables_processed,
                duration_seconds=duration,
            )
            
        except Exception as e:
            logger.error(f"DOCX protection failed: {e}", exc_info=True)
            return DocxProtectionResult(
                success=False,
                error_message=str(e),
            )
        finally:
            if gc_was_enabled:
                gc.enable()
    
    def _replace_in_paragraph(self, paragraph, value_to_token: dict[str, str]) -> int:
        """
        Replace values in a paragraph while preserving formatting.
        
        Returns count of replacements made.
        """
        count = 0
        
        # Get full text
        full_text = paragraph.text
        if not full_text:
            return 0
        
        # Check if any values exist in this paragraph
        values_in_para = [v for v in value_to_token.keys() if v in full_text]
        if not values_in_para:
            return 0
        
        # Replace in each run (preserves formatting)
        for run in paragraph.runs:
            text = run.text
            if not text:
                continue
            
            new_text = text
            for value, token in value_to_token.items():
                if value in new_text:
                    new_text = new_text.replace(value, token)
                    count += text.count(value)
            
            if new_text != text:
                run.text = new_text
        
        return count
