"""
Excel file protection service.

Replaces selected column values with tokens while preserving file structure.
"""

from __future__ import annotations

import logging
import os
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING, Callable

import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from sandiraksa.protection.tokenizer import Tokenizer, TokenizerFactory

if TYPE_CHECKING:
    pass

logger = logging.getLogger(__name__)


@dataclass
class ColumnConfig:
    """Configuration for a column to protect."""
    
    sheet_name: str
    column_index: int  # 1-based
    column_name: str
    entity_type: str = "PII"  # Default entity type
    # "full_cell" (replace whole cell with one token) or "pii_only" (detect and
    # tokenize only PII substrings inside the cell, keeping the rest intact).
    protection_mode: str = "full_cell"


@dataclass
class ProtectionResult:
    """Result of protecting a file."""
    
    success: bool
    output_path: str | None = None
    error_message: str | None = None
    cells_protected: int = 0
    rows_processed: int = 0
    tokens_created: int = 0
    duration_seconds: float = 0.0


@dataclass
class ProtectionProgress:
    """Progress information during protection."""
    
    current_sheet: str = ""
    current_row: int = 0
    total_rows: int = 0
    cells_protected: int = 0


class ExcelProtector:
    """
    Protects Excel files by replacing selected column values with tokens.
    
    Features:
    - Preserves original file structure and formatting
    - Processes only selected columns
    - Uses consistent tokenization (same value = same token)
    - Reports progress for UI feedback
    """
    
    def __init__(
        self,
        project_id: str,
        tokenizer: Tokenizer | None = None,
    ) -> None:
        """
        Initialize Excel protector.
        
        Args:
            project_id: Project ID for token scoping.
            tokenizer: Optional tokenizer instance. If None, creates one.
        """
        self._project_id = project_id
        self._tokenizer = tokenizer or TokenizerFactory.get_tokenizer(project_id)
    
    def protect_file(
        self,
        input_path: str | Path,
        output_path: str | Path,
        columns: list[ColumnConfig],
        progress_callback: Callable[[ProtectionProgress], None] | None = None,
    ) -> ProtectionResult:
        """
        Protect an Excel file by tokenizing selected columns.
        
        Args:
            input_path: Path to input Excel file.
            output_path: Path where protected file will be saved.
            columns: List of columns to protect.
            progress_callback: Optional callback for progress updates.
            
        Returns:
            ProtectionResult with status and statistics.
        """
        input_path = Path(input_path)
        output_path = Path(output_path)
        
        start_time = datetime.now()
        progress = ProtectionProgress()
        cells_protected = 0
        rows_processed = 0
        
        try:
            # Load workbook with data_only=False to preserve formulas
            logger.info(f"Loading workbook: {input_path}")
            wb = load_workbook(input_path, data_only=False)
            
            # Group columns by sheet
            columns_by_sheet: dict[str, list[ColumnConfig]] = {}
            for col in columns:
                if col.sheet_name not in columns_by_sheet:
                    columns_by_sheet[col.sheet_name] = []
                columns_by_sheet[col.sheet_name].append(col)
            
            # Process each sheet
            for sheet_name, sheet_columns in columns_by_sheet.items():
                if sheet_name not in wb.sheetnames:
                    logger.warning(f"Sheet not found: {sheet_name}")
                    continue
                
                ws = wb[sheet_name]
                progress.current_sheet = sheet_name
                progress.total_rows = ws.max_row - 1  # Exclude header
                
                # Process each column
                for col_config in sheet_columns:
                    protected = self._protect_column(
                        ws,
                        col_config,
                        progress,
                        progress_callback,
                    )
                    cells_protected += protected
                
                rows_processed += progress.total_rows
            
            # Ensure output directory exists
            output_path.parent.mkdir(parents=True, exist_ok=True)
            
            # Save protected workbook
            logger.info(f"Saving protected file: {output_path}")
            wb.save(output_path)
            wb.close()
            
            duration = (datetime.now() - start_time).total_seconds()
            
            logger.info(
                f"Protection complete: {cells_protected} cells protected "
                f"in {duration:.2f}s"
            )
            
            return ProtectionResult(
                success=True,
                output_path=str(output_path),
                cells_protected=cells_protected,
                rows_processed=rows_processed,
                tokens_created=self._tokenizer.get_mapping_count(),
                duration_seconds=duration,
            )
            
        except Exception as e:
            logger.error(f"Protection failed: {e}", exc_info=True)
            return ProtectionResult(
                success=False,
                error_message=str(e),
            )
    
    def _protect_column(
        self,
        ws: Worksheet,
        col_config: ColumnConfig,
        progress: ProtectionProgress,
        progress_callback: Callable[[ProtectionProgress], None] | None,
    ) -> int:
        """
        Protect a single column in a worksheet.
        
        Returns number of cells protected.
        """
        cells_protected = 0
        col_idx = col_config.column_index
        entity_type = col_config.entity_type
        pii_only = col_config.protection_mode == "pii_only"

        # For PII-only mode, build the detection engine once per column.
        engine = self._build_engine() if pii_only else None

        # Start from row 2 (skip header)
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=col_idx)
            value = cell.value
            
            # Skip empty cells
            if value is None or (isinstance(value, str) and not value.strip()):
                continue
            
            # Convert to string for tokenization
            str_value = str(value).strip()

            if pii_only and engine is not None:
                # Tokenize ONLY the detected PII substrings inside the cell,
                # leaving surrounding free text / JSON structure intact.
                new_value, hits = self._tokenize_pii_in_text(engine, str_value)
                if hits > 0:
                    cell.value = new_value
                    cells_protected += 1
            else:
                # Replace the whole cell value with one token.
                token = self._tokenizer.get_or_create_token(entity_type, str_value)
                cell.value = token
                cells_protected += 1
            
            # Update progress
            progress.current_row = row - 1
            progress.cells_protected = cells_protected
            
            # Report progress every 100 rows
            if progress_callback and (row % 100 == 0):
                progress_callback(progress)
        
        # Final progress update
        if progress_callback:
            progress_callback(progress)
        
        return cells_protected

    def _build_engine(self):
        """Build a detection engine with the standard recognizer set."""
        from sandiraksa.detection.engine import DetectionEngine
        from sandiraksa.detection.presidio_engine import RegexRecognizer
        from sandiraksa.detection.recognizers import (
            NIKRecognizer,
            NPWPRecognizer,
            KKRecognizer,
            IndonesianPhoneRecognizer,
            BPJSRecognizerLegacy,
        )
        from sandiraksa.detection.recognizers.id_dob import DateOfBirthRecognizer
        from sandiraksa.detection.recognizers.id_person import (
            IndonesianPersonRecognizer,
        )

        engine = DetectionEngine()
        engine.registry.register(RegexRecognizer())
        engine.registry.register(NIKRecognizer())
        engine.registry.register(NPWPRecognizer())
        engine.registry.register(KKRecognizer())
        engine.registry.register(IndonesianPhoneRecognizer())
        engine.registry.register(BPJSRecognizerLegacy())
        engine.registry.register(DateOfBirthRecognizer())
        engine.registry.register(IndonesianPersonRecognizer())

        # NLP (Presidio/spaCy) for names inside free text / JSON. This is what
        # lets "Budi Sanjaya" in a narrative cell be detected (the context-only
        # IndonesianPersonRecognizer cannot see unlabeled names). Fail-safe: if
        # Presidio/spaCy is unavailable, skip it and keep the rest working.
        try:
            from sandiraksa.detection.presidio_engine import PresidioRecognizer

            engine.registry.register(
                PresidioRecognizer(entities=["PERSON"], language="en")
            )
        except Exception as e:
            logger.warning(f"Presidio NER unavailable, names in free text may be missed: {e}")

        engine.initialize()
        return engine

    def _tokenize_pii_in_text(self, engine, text: str) -> tuple[str, int]:
        """
        Detect PII inside ``text`` and replace ONLY those substrings with tokens.

        Returns (new_text, number_of_substitutions). Replacements are applied
        right-to-left so earlier offsets stay valid.
        """
        from uuid import uuid4
        from sandiraksa.detection.context import DetectionContext, DetectionConfig

        config = DetectionConfig(
            enabled_entity_types={
                "EMAIL_ADDRESS", "PHONE_NUMBER", "CREDIT_CARD",
                "IP_ADDRESS", "URL", "IBAN_CODE",
                "ID_NIK", "ID_NPWP", "ID_KK", "ID_PHONE", "ID_BPJS",
                "PERSON", "DATE_OF_BIRTH",
            },
            min_confidence=0.5,
        )
        context = DetectionContext(
            operation_id=str(uuid4()),
            file_id=str(uuid4()),
            project_id=self._project_id,
            config=config,
        )

        results = list(engine.analyze_text(text, context))

        # Chain global custom patterns (e.g. Medical Record formats). The engine
        # does not run these, so add them here as high-confidence hits.
        results.extend(self._custom_pattern_hits(text))

        # Trim NLP-detected PERSON spans that over-extend into label/noise words
        # (e.g. "Pasien Budi S" -> "Budi S"). Drops spans that trim to nothing.
        results = self._refine_person_spans(results, text)

        if not results:
            return text, 0

        # Greedy non-overlap: prefer higher score; on conflict skip the loser.
        chosen: list = []
        occupied: list[tuple[int, int]] = []
        for r in sorted(results, key=lambda r: (-r.score, r.start)):
            if any(not (r.end <= s or r.start >= e) for s, e in occupied):
                continue
            occupied.append((r.start, r.end))
            chosen.append(r)

        new_text = text
        for r in sorted(chosen, key=lambda r: r.start, reverse=True):
            token = self._tokenizer.get_or_create_token(r.entity_type, r.text)
            new_text = new_text[: r.start] + token + new_text[r.end :]

        return new_text, len(chosen)

    def _refine_person_spans(self, results: list, text: str) -> list:
        """
        Tighten NLP PERSON spans; drop those that trim away to non-names.

        Non-PERSON results pass through unchanged.
        """
        try:
            from sandiraksa.detection.recognizers.person_filter import (
                trim_person_span,
            )
        except Exception:
            return results

        refined = []
        for r in results:
            if getattr(r, "entity_type", None) != "PERSON":
                refined.append(r)
                continue
            trimmed = trim_person_span(text, r.start, r.end)
            if trimmed is None:
                continue  # trimmed to nothing -> not a name, drop
            clean, ns, ne = trimmed
            r.text = clean
            r.start = ns
            r.end = ne
            refined.append(r)
        return refined

    def _custom_pattern_hits(self, text: str):
        """
        Return global custom-pattern matches as DetectionResult-like objects.

        The detection engine does not evaluate user-defined custom patterns
        (e.g. hospital Medical Record formats), so we chain them here. Given a
        high score so they win overlap resolution against generic matches.
        """
        from sandiraksa.detection.context import DetectionResult

        hits: list[DetectionResult] = []
        try:
            from sandiraksa.detection.custom_patterns import find_custom_matches

            for m in find_custom_matches(text):
                hits.append(DetectionResult(
                    entity_type=m.label,
                    start=m.start,
                    end=m.end,
                    text=m.value,
                    score=0.95,
                    recognizer_name="custom_pattern",
                ))
        except Exception as e:
            logger.warning(f"Custom pattern detection failed: {e}")
        return hits


def generate_output_path(
    input_path: str | Path,
    output_dir: str | Path | None = None,
    suffix: str = "_protected",
) -> Path:
    """
    Generate output path for protected file.
    
    Args:
        input_path: Original file path.
        output_dir: Output directory. If None, uses same dir as input.
        suffix: Suffix to add before extension.
        
    Returns:
        Path for output file.
    """
    input_path = Path(input_path)
    
    if output_dir is None:
        output_dir = input_path.parent
    else:
        output_dir = Path(output_dir)
    
    stem = input_path.stem
    ext = input_path.suffix
    
    # Generate unique name if file exists
    output_path = output_dir / f"{stem}{suffix}{ext}"
    counter = 1
    
    while output_path.exists():
        output_path = output_dir / f"{stem}{suffix}_{counter}{ext}"
        counter += 1
    
    return output_path


def parse_selected_columns(
    selected_columns: dict[str, list[tuple[int, str]]],
    entity_type_map: dict[str, str] | None = None,
) -> list[ColumnConfig]:
    """
    Parse selected columns dict to ColumnConfig list.
    
    Args:
        selected_columns: Dict of {sheet_name: [(col_idx, col_name), ...]}
        entity_type_map: Optional mapping of column names to entity types.
        
    Returns:
        List of ColumnConfig objects.
    """
    configs = []
    entity_type_map = entity_type_map or {}
    
    # Default entity type mapping based on column name keywords
    default_entity_types = {
        "nama": "PERSON",
        "name": "PERSON",
        "email": "EMAIL",
        "e-mail": "EMAIL",
        "phone": "PHONE_NUMBER",
        "telepon": "PHONE_NUMBER",
        "telp": "PHONE_NUMBER",
        "hp": "PHONE_NUMBER",
        "handphone": "PHONE_NUMBER",
        "nik": "ID_NIK",
        "ktp": "ID_NIK",
        "npwp": "ID_NPWP",
        "kk": "ID_KK",
        "alamat": "ADDRESS",
        "address": "ADDRESS",
        "rekam medis": "MEDICAL_RECORD",
        "medical record": "MEDICAL_RECORD",
        "mrn": "MEDICAL_RECORD",
        "no rm": "MEDICAL_RECORD",
        "norm": "MEDICAL_RECORD",
    }
    
    for sheet_name, columns in selected_columns.items():
        for col_idx, col_name in columns:
            # Determine entity type
            entity_type = entity_type_map.get(col_name, "PII")
            
            # Try to infer from column name
            if entity_type == "PII":
                col_name_lower = col_name.lower()
                for keyword, etype in default_entity_types.items():
                    if keyword in col_name_lower:
                        entity_type = etype
                        break
            
            configs.append(ColumnConfig(
                sheet_name=sheet_name,
                column_index=col_idx,
                column_name=col_name,
                entity_type=entity_type,
            ))
    
    return configs


__all__ = [
    "ColumnConfig",
    "ExcelProtector",
    "ProtectionProgress",
    "ProtectionResult",
    "generate_output_path",
    "parse_selected_columns",
]
