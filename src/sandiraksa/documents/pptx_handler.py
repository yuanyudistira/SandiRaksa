"""
PPTX Document Handler.

Provides PowerPoint presentation processing with:
- Slide inventory and navigation
- Shape text extraction (text boxes, placeholders)
- Table text extraction
- Notes extraction
- Style-preserving modifications
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING, Any, Callable, Iterator
from zipfile import ZipFile

from sandiraksa.detection.context import TextSegment
from sandiraksa.domain.finding import DocumentLocation
from sandiraksa.domain.file_record import FileFormat

if TYPE_CHECKING:
    from pptx import Presentation
    from pptx.shapes.base import BaseShape
    from pptx.slide import Slide

logger = logging.getLogger(__name__)


class PPTXHandlerError(Exception):
    """Base exception for PPTX handler errors."""

    pass


class PresentationOpenError(PPTXHandlerError):
    """Failed to open presentation."""

    pass


class SlideModificationError(PPTXHandlerError):
    """Failed to modify slide."""

    pass


@dataclass
class PPTXMetadata:
    """Metadata about a PPTX file."""

    file_path: Path
    file_size: int
    slide_count: int
    has_notes: bool
    has_comments: bool
    has_tables: bool
    created: str | None = None
    modified: str | None = None
    creator: str | None = None
    title: str | None = None
    subject: str | None = None


@dataclass
class ShapeInfo:
    """Information about a shape in a slide."""

    slide_number: int  # 1-indexed
    shape_id: int
    shape_type: str  # "text_box", "placeholder", "table", "chart", etc.
    text: str
    has_text_frame: bool = True
    paragraph_count: int = 0
    
    # Position
    left: int = 0
    top: int = 0
    width: int = 0
    height: int = 0

    def to_location(self, file_id: str = "") -> DocumentLocation:
        """Convert to DocumentLocation."""
        return DocumentLocation(
            file_id=file_id,
            component_type="shape",
            slide_number=self.slide_number,
            shape_id=str(self.shape_id),
        )


@dataclass
class TableCellInfo:
    """Information about a table cell."""

    slide_number: int
    shape_id: int
    row_index: int
    column_index: int
    text: str

    def to_location(self, file_id: str = "") -> DocumentLocation:
        """Convert to DocumentLocation."""
        return DocumentLocation(
            file_id=file_id,
            component_type="table_cell",
            slide_number=self.slide_number,
            shape_id=str(self.shape_id),
        )


@dataclass
class SlideNotesInfo:
    """Information about slide notes."""

    slide_number: int
    text: str

    def to_location(self, file_id: str = "") -> DocumentLocation:
        """Convert to DocumentLocation."""
        return DocumentLocation(
            file_id=file_id,
            component_type="notes",
            slide_number=self.slide_number,
        )


@dataclass
class SlideInventory:
    """Inventory of a single slide."""

    slide_number: int
    shape_count: int
    text_shape_count: int
    table_count: int
    has_notes: bool
    title: str | None = None


@dataclass
class PresentationInventory:
    """Complete inventory of a presentation."""

    file_path: Path
    metadata: PPTXMetadata
    slides: list[SlideInventory]
    total_shapes: int
    total_tables: int


class PPTXReader:
    """
    Reads PPTX files with comprehensive text extraction.

    Handles:
    - Slide shapes (text boxes, placeholders)
    - Tables
    - Notes
    """

    def __init__(self, file_path: Path | str) -> None:
        """
        Initialize PPTX reader.

        Args:
            file_path: Path to PPTX file.
        """
        self._file_path = Path(file_path)
        self._presentation: Presentation | None = None
        self._inventory: PresentationInventory | None = None

    def __enter__(self) -> PPTXReader:
        """Context manager entry."""
        self.open()
        return self

    def __exit__(self, *args: Any) -> None:
        """Context manager exit."""
        self.close()

    def open(self) -> None:
        """Open the presentation."""
        try:
            from pptx import Presentation

            self._presentation = Presentation(self._file_path)
            logger.debug(f"Opened presentation: {self._file_path}")
        except Exception as e:
            raise PresentationOpenError(f"Failed to open {self._file_path}: {e}") from e

    def close(self) -> None:
        """Close the presentation."""
        self._presentation = None

    @property
    def presentation(self) -> Presentation:
        """Get the presentation, opening if needed."""
        if self._presentation is None:
            self.open()
        assert self._presentation is not None
        return self._presentation

    @property
    def slide_count(self) -> int:
        """Get number of slides."""
        return len(self.presentation.slides)

    def get_inventory(self) -> PresentationInventory:
        """
        Build inventory of the presentation.

        Returns:
            PresentationInventory with slide details.
        """
        if self._inventory is not None:
            return self._inventory

        prs = self.presentation
        slides: list[SlideInventory] = []
        total_shapes = 0
        total_tables = 0
        has_notes = False
        has_tables = False

        for slide_num, slide in enumerate(prs.slides, 1):
            shape_count = len(slide.shapes)
            text_shape_count = 0
            table_count = 0
            slide_has_notes = False
            title = None

            for shape in slide.shapes:
                if shape.has_text_frame:
                    text_shape_count += 1
                if shape.has_table:
                    table_count += 1
                    has_tables = True
                
                # Get title
                if shape.is_placeholder:
                    if hasattr(shape, "placeholder_format"):
                        from pptx.enum.shapes import PP_PLACEHOLDER
                        if shape.placeholder_format.type == PP_PLACEHOLDER.TITLE:
                            if shape.has_text_frame:
                                title = shape.text_frame.text

            # Check for notes
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                if notes_slide.notes_text_frame:
                    if notes_slide.notes_text_frame.text.strip():
                        slide_has_notes = True
                        has_notes = True

            total_shapes += shape_count
            total_tables += table_count

            slides.append(
                SlideInventory(
                    slide_number=slide_num,
                    shape_count=shape_count,
                    text_shape_count=text_shape_count,
                    table_count=table_count,
                    has_notes=slide_has_notes,
                    title=title,
                )
            )

        # Build metadata
        core_props = prs.core_properties
        metadata = PPTXMetadata(
            file_path=self._file_path,
            file_size=self._file_path.stat().st_size,
            slide_count=len(slides),
            has_notes=has_notes,
            has_comments=self._check_has_comments(),
            has_tables=has_tables,
            created=str(core_props.created) if core_props.created else None,
            modified=str(core_props.modified) if core_props.modified else None,
            creator=core_props.author,
            title=core_props.title,
            subject=core_props.subject,
        )

        self._inventory = PresentationInventory(
            file_path=self._file_path,
            metadata=metadata,
            slides=slides,
            total_shapes=total_shapes,
            total_tables=total_tables,
        )

        return self._inventory

    def _check_has_comments(self) -> bool:
        """Check if presentation has comments."""
        try:
            with ZipFile(self._file_path, "r") as zf:
                return any("comments" in name for name in zf.namelist())
        except Exception:
            return False

    def iter_shapes(
        self,
        slide_number: int | None = None,
        include_tables: bool = True,
    ) -> Iterator[ShapeInfo]:
        """
        Iterate over text shapes.

        Args:
            slide_number: Specific slide (1-indexed) or None for all.
            include_tables: Include table shapes.

        Yields:
            ShapeInfo for each shape with text.
        """
        prs = self.presentation

        if slide_number:
            slides = [prs.slides[slide_number - 1]]
            start_num = slide_number
        else:
            slides = prs.slides
            start_num = 1

        for idx, slide in enumerate(slides, start_num):
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text
                    if text and text.strip():
                        # Determine shape type
                        shape_type = "text_box"
                        if shape.is_placeholder:
                            shape_type = "placeholder"

                        yield ShapeInfo(
                            slide_number=idx,
                            shape_id=shape.shape_id,
                            shape_type=shape_type,
                            text=text,
                            has_text_frame=True,
                            paragraph_count=len(shape.text_frame.paragraphs),
                            left=shape.left or 0,
                            top=shape.top or 0,
                            width=shape.width or 0,
                            height=shape.height or 0,
                        )

                # Handle tables separately if needed
                if include_tables and shape.has_table:
                    yield ShapeInfo(
                        slide_number=idx,
                        shape_id=shape.shape_id,
                        shape_type="table",
                        text="[TABLE]",  # Placeholder - cells yielded separately
                        has_text_frame=False,
                    )

    def iter_table_cells(
        self,
        slide_number: int | None = None,
    ) -> Iterator[TableCellInfo]:
        """
        Iterate over table cells.

        Args:
            slide_number: Specific slide or None for all.

        Yields:
            TableCellInfo for each cell with text.
        """
        prs = self.presentation

        if slide_number:
            slides = [(slide_number, prs.slides[slide_number - 1])]
        else:
            slides = enumerate(prs.slides, 1)

        for slide_num, slide in slides:
            for shape in slide.shapes:
                if shape.has_table:
                    table = shape.table
                    for row_idx, row in enumerate(table.rows):
                        for col_idx, cell in enumerate(row.cells):
                            text = cell.text
                            if text and text.strip():
                                yield TableCellInfo(
                                    slide_number=slide_num,
                                    shape_id=shape.shape_id,
                                    row_index=row_idx,
                                    column_index=col_idx,
                                    text=text,
                                )

    def iter_notes(self) -> Iterator[SlideNotesInfo]:
        """
        Iterate over slide notes.

        Yields:
            SlideNotesInfo for each slide with notes.
        """
        prs = self.presentation

        for slide_num, slide in enumerate(prs.slides, 1):
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                if notes_slide.notes_text_frame:
                    text = notes_slide.notes_text_frame.text
                    if text and text.strip():
                        yield SlideNotesInfo(
                            slide_number=slide_num,
                            text=text,
                        )

    def iter_text_segments(self) -> Iterator[TextSegment]:
        """
        Iterate all text as TextSegments for detection.

        Yields:
            TextSegment for each text element.
        """
        # Shapes
        for shape in self.iter_shapes(include_tables=False):
            yield TextSegment(
                text=shape.text,
                location=shape.to_location(),
            )

        # Table cells
        for cell in self.iter_table_cells():
            yield TextSegment(
                text=cell.text,
                location=cell.to_location(),
            )

        # Notes
        for notes in self.iter_notes():
            yield TextSegment(
                text=notes.text,
                location=notes.to_location(),
            )


class PPTXWriter:
    """
    Modifies PPTX files with style preservation.
    """

    def __init__(
        self,
        source_path: Path | str,
        output_path: Path | str | None = None,
    ) -> None:
        """
        Initialize PPTX writer.

        Args:
            source_path: Source PPTX file.
            output_path: Output path.
        """
        self._source_path = Path(source_path)
        self._output_path = Path(output_path) if output_path else self._source_path
        self._presentation: Presentation | None = None
        self._modifications: int = 0

    def __enter__(self) -> PPTXWriter:
        """Context manager entry."""
        self.open()
        return self

    def __exit__(self, *args: Any) -> None:
        """Context manager exit."""
        self.save()

    def open(self) -> None:
        """Open presentation for modification."""
        try:
            from pptx import Presentation

            self._presentation = Presentation(self._source_path)
            logger.debug(f"Opened presentation for writing: {self._source_path}")
        except Exception as e:
            raise PresentationOpenError(f"Failed to open {self._source_path}: {e}") from e

    def save(self) -> None:
        """Save changes."""
        if self._presentation is not None:
            self._presentation.save(self._output_path)
            logger.info(
                f"Saved presentation: {self._output_path} "
                f"({self._modifications} modifications)"
            )

    @property
    def presentation(self) -> Presentation:
        """Get the presentation."""
        if self._presentation is None:
            self.open()
        assert self._presentation is not None
        return self._presentation

    def replace_in_shape(
        self,
        slide_number: int,
        shape_id: int,
        old_text: str,
        new_text: str,
    ) -> bool:
        """
        Replace text in a shape.

        Args:
            slide_number: 1-indexed slide number.
            shape_id: Shape ID.
            old_text: Text to find.
            new_text: Replacement text.

        Returns:
            True if replacement was made.
        """
        prs = self.presentation

        if slide_number < 1 or slide_number > len(prs.slides):
            return False

        slide = prs.slides[slide_number - 1]

        for shape in slide.shapes:
            if shape.shape_id == shape_id and shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if old_text in run.text:
                            run.text = run.text.replace(old_text, new_text)
                            self._modifications += 1
                            return True

                # If not found in runs, try full text replacement
                full_text = shape.text_frame.text
                if old_text in full_text:
                    # Simple replacement in first paragraph/run
                    if shape.text_frame.paragraphs:
                        para = shape.text_frame.paragraphs[0]
                        if para.runs:
                            para.runs[0].text = full_text.replace(old_text, new_text)
                        else:
                            para.text = full_text.replace(old_text, new_text)
                        self._modifications += 1
                        return True

        return False

    def replace_in_table_cell(
        self,
        slide_number: int,
        shape_id: int,
        row_index: int,
        column_index: int,
        old_text: str,
        new_text: str,
    ) -> bool:
        """
        Replace text in a table cell.

        Args:
            slide_number: 1-indexed slide number.
            shape_id: Table shape ID.
            row_index: Row index.
            column_index: Column index.
            old_text: Text to find.
            new_text: Replacement text.

        Returns:
            True if replacement was made.
        """
        prs = self.presentation

        if slide_number < 1 or slide_number > len(prs.slides):
            return False

        slide = prs.slides[slide_number - 1]

        for shape in slide.shapes:
            if shape.shape_id == shape_id and shape.has_table:
                table = shape.table
                if row_index < len(table.rows):
                    row = table.rows[row_index]
                    if column_index < len(row.cells):
                        cell = row.cells[column_index]
                        if old_text in cell.text:
                            # Replace in cell
                            for para in cell.text_frame.paragraphs:
                                for run in para.runs:
                                    if old_text in run.text:
                                        run.text = run.text.replace(old_text, new_text)
                                        self._modifications += 1
                                        return True
        return False

    def replace_in_notes(
        self,
        slide_number: int,
        old_text: str,
        new_text: str,
    ) -> bool:
        """
        Replace text in slide notes.

        Args:
            slide_number: 1-indexed slide number.
            old_text: Text to find.
            new_text: Replacement text.

        Returns:
            True if replacement was made.
        """
        prs = self.presentation

        if slide_number < 1 or slide_number > len(prs.slides):
            return False

        slide = prs.slides[slide_number - 1]

        if slide.has_notes_slide:
            notes_slide = slide.notes_slide
            if notes_slide.notes_text_frame:
                for para in notes_slide.notes_text_frame.paragraphs:
                    for run in para.runs:
                        if old_text in run.text:
                            run.text = run.text.replace(old_text, new_text)
                            self._modifications += 1
                            return True
        return False

    def apply_treatments(
        self,
        treatments: list[dict],
        progress_callback: Callable[[int, int], None] | None = None,
    ) -> int:
        """
        Apply multiple treatments.

        Args:
            treatments: List of treatment dicts with:
                - slide_number
                - old_text
                - new_text
                - component_type (shape, table_cell, notes)
                - shape_id (for shapes and tables)
                - row_index, column_index (for table cells)
            progress_callback: Called with (current, total).

        Returns:
            Number of successful modifications.
        """
        total = len(treatments)
        modified = 0

        for idx, treatment in enumerate(treatments):
            try:
                component_type = treatment.get("component_type", "shape")
                success = False

                if component_type == "shape":
                    success = self.replace_in_shape(
                        slide_number=treatment["slide_number"],
                        shape_id=treatment["shape_id"],
                        old_text=treatment["old_text"],
                        new_text=treatment["new_text"],
                    )
                elif component_type == "table_cell":
                    success = self.replace_in_table_cell(
                        slide_number=treatment["slide_number"],
                        shape_id=treatment["shape_id"],
                        row_index=treatment["row_index"],
                        column_index=treatment["column_index"],
                        old_text=treatment["old_text"],
                        new_text=treatment["new_text"],
                    )
                elif component_type == "notes":
                    success = self.replace_in_notes(
                        slide_number=treatment["slide_number"],
                        old_text=treatment["old_text"],
                        new_text=treatment["new_text"],
                    )

                if success:
                    modified += 1

            except Exception as e:
                logger.warning(f"Failed to apply treatment: {e}")

            if progress_callback:
                progress_callback(idx + 1, total)

        return modified


class PPTXHandler:
    """
    High-level PPTX handler combining reader and writer.
    """

    def __init__(self, file_path: Path | str) -> None:
        """
        Initialize handler.

        Args:
            file_path: Path to PPTX file.
        """
        self._file_path = Path(file_path)

    @property
    def file_format(self) -> FileFormat:
        """Get file format."""
        return FileFormat.PPTX

    def get_metadata(self) -> PPTXMetadata:
        """Get presentation metadata."""
        with PPTXReader(self._file_path) as reader:
            inventory = reader.get_inventory()
            return inventory.metadata

    def get_inventory(self) -> PresentationInventory:
        """Get full presentation inventory."""
        with PPTXReader(self._file_path) as reader:
            return reader.get_inventory()

    def iter_text_segments(self) -> Iterator[TextSegment]:
        """
        Iterate all text segments for detection.

        Yields:
            TextSegment for each text element.
        """
        with PPTXReader(self._file_path) as reader:
            yield from reader.iter_text_segments()

    def iter_shapes(
        self,
        slide_number: int | None = None,
    ) -> Iterator[ShapeInfo]:
        """
        Iterate shapes in presentation.

        Yields:
            ShapeInfo for each shape.
        """
        with PPTXReader(self._file_path) as reader:
            yield from reader.iter_shapes(slide_number=slide_number)

    def apply_treatments(
        self,
        treatments: list[dict],
        output_path: Path | str | None = None,
        progress_callback: Callable[[int, int], None] | None = None,
    ) -> tuple[Path, int]:
        """
        Apply treatments and save.

        Returns:
            Tuple of (output_path, modifications_count).
        """
        output = Path(output_path) if output_path else self._get_default_output_path()

        with PPTXWriter(self._file_path, output) as writer:
            modified = writer.apply_treatments(treatments, progress_callback)

        return output, modified

    def _get_default_output_path(self) -> Path:
        """Generate default output path."""
        stem = self._file_path.stem
        suffix = self._file_path.suffix
        return self._file_path.parent / f"{stem}_protected{suffix}"

    def clear_metadata(
        self,
        output_path: Path | str | None = None,
        clear_author: bool = True,
        clear_title: bool = False,
    ) -> Path:
        """
        Clear sensitive metadata.

        Returns:
            Output file path.
        """
        output = Path(output_path) if output_path else self._get_default_output_path()

        from pptx import Presentation

        prs = Presentation(self._file_path)

        if clear_author:
            prs.core_properties.author = None
            prs.core_properties.last_modified_by = None

        if clear_title:
            prs.core_properties.title = None
            prs.core_properties.subject = None

        prs.save(output)
        logger.info(f"Cleared metadata and saved to: {output}")
        return output


__all__ = [
    # Errors
    "PPTXHandlerError",
    "PresentationOpenError",
    "SlideModificationError",
    # Data classes
    "PPTXMetadata",
    "ShapeInfo",
    "TableCellInfo",
    "SlideNotesInfo",
    "SlideInventory",
    "PresentationInventory",
    # Classes
    "PPTXReader",
    "PPTXWriter",
    "PPTXHandler",
]
