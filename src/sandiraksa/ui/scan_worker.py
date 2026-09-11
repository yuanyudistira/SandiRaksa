"""
Background scan worker.

Runs file reading and PII detection off the GUI thread so the Qt event loop
is never blocked or re-entered during a scan. This removes the
processEvents()/self-rescheduling-QTimer re-entrancy that (combined with a
disabled garbage collector) caused intermittent Windows heap corruption
(STATUS_HEAP_CORRUPTION, 0xC0000374) when processing large Excel files.

Design notes for large files (tens of thousands of rows x many columns):
  * The detection engine and the DetectionConfig/DetectionContext objects are
    built ONCE per worker run and reused across every cell, instead of being
    allocated per cell. This slashes short-lived allocations.
  * The worker periodically triggers gc.collect() so transient objects
    (openpyxl rows, result lists, finding dicts) are reclaimed steadily rather
    than piling up.
  * Cancellation is cooperative via a thread-safe flag checked in the row loop.
"""

from __future__ import annotations

import gc
from pathlib import Path
from uuid import uuid4

from PySide6.QtCore import QObject, Signal


# Entity types detected for structured/typed cells.
_STRUCTURED_ENTITIES = {
    "EMAIL_ADDRESS", "PHONE_NUMBER", "CREDIT_CARD",
    "IP_ADDRESS", "URL", "IBAN_CODE",
    "ID_NIK", "ID_NPWP", "ID_KK", "ID_PHONE", "ID_BPJS",
}

# Free-text / narrative columns additionally look for names and birth dates.
_PII_ONLY_ENTITIES = _STRUCTURED_ENTITIES | {"PERSON", "DATE_OF_BIRTH"}

# Full-content (non-column) scan entity set.
_FULL_CONTENT_ENTITIES = _STRUCTURED_ENTITIES | {"PERSON", "DATE_OF_BIRTH"}

# Types marked as-is (no regex detection needed) when a column is typed.
_DIRECT_TYPES = ("PERSON", "ADDRESS", "MEDICAL_RECORD", "MEDICAL_INFO", "DATE_OF_BIRTH")

# Findings caps (kept identical to the previous inline behavior).
_MAX_PER_TYPE = 1000
_MAX_TOTAL = 10000

# Run gc.collect() every N processed cells during heavy column scans.
_GC_EVERY_CELLS = 5000


class ScanWorker(QObject):
    """
    Scans a list of files for PII on a background thread.

    Move an instance to a QThread and connect the thread's ``started`` signal
    to :meth:`run`. Results and progress are delivered via queued signals, so
    all GUI updates happen safely on the main thread.
    """

    # file_id
    file_started = Signal(str)
    # file_id, findings_count
    file_finished = Signal(str, int)
    # file_id
    file_error = Signal(str)
    # current, total, message
    progress = Signal(int, int, str)
    # full list of finding dicts (already tagged with ids/file info)
    all_finished = Signal(list)

    def __init__(
        self,
        files: list[dict],
        project_id: str | None,
        parent: QObject | None = None,
    ) -> None:
        super().__init__(parent)
        self._files = files
        self._project_id = project_id or str(uuid4())
        self._cancelled = False
        self._engine = None
        # Reusable detection contexts (built lazily, reused across all cells).
        self._ctx_structured = None
        self._ctx_pii_only = None
        self._ctx_full = None

    # ------------------------------------------------------------------ API

    def cancel(self) -> None:
        """Request cooperative cancellation (thread-safe: single bool write)."""
        self._cancelled = True

    def run(self) -> None:
        """Entry point executed on the worker thread."""
        all_findings: list[dict] = []
        try:
            self._ensure_engine()
            total = len(self._files)
            for index, file_info in enumerate(self._files):
                if self._cancelled:
                    break

                file_id = file_info.get("id", str(uuid4()))
                file_path = file_info.get("path", "")
                name = file_info.get("name", "Unknown")

                self.file_started.emit(file_id)
                self.progress.emit(index, total, name)

                try:
                    if file_path and Path(file_path).exists():
                        raw = self._scan_file(
                            file_path, file_info.get("selected_columns")
                        )
                        selected = self._finalize_findings(raw, file_id, name)
                        all_findings.extend(selected)
                        self.file_finished.emit(file_id, len(selected))
                    else:
                        self.file_finished.emit(file_id, 0)
                except Exception as exc:  # pragma: no cover - defensive
                    print(f"ScanWorker: error scanning {file_path}: {exc}")
                    import traceback

                    traceback.print_exc()
                    self.file_error.emit(file_id)

                # Reclaim transient per-file allocations promptly.
                gc.collect()
        except Exception as exc:  # pragma: no cover - defensive
            print(f"ScanWorker: fatal error: {exc}")
            import traceback

            traceback.print_exc()
        finally:
            self.all_finished.emit(all_findings)

    # ------------------------------------------------------- engine / context

    def _ensure_engine(self) -> None:
        """Build the detection engine and reusable contexts once."""
        if self._engine is not None:
            return

        from sandiraksa.detection.presidio_engine import RegexRecognizer
        from sandiraksa.detection.engine import DetectionEngine
        from sandiraksa.detection.recognizers import (
            NIKRecognizer,
            NPWPRecognizer,
            KKRecognizer,
            IndonesianPhoneRecognizer,
        )
        from sandiraksa.detection.recognizers.id_dob import DateOfBirthRecognizer
        from sandiraksa.detection.recognizers.id_person import (
            IndonesianPersonRecognizer,
        )
        from sandiraksa.detection.recognizers.id_bpjs_legacy import (
            BPJSRecognizerLegacy,
        )

        engine = DetectionEngine()
        engine.registry.register(RegexRecognizer())
        engine.registry.register(NIKRecognizer())
        engine.registry.register(NPWPRecognizer())
        engine.registry.register(KKRecognizer())
        engine.registry.register(IndonesianPhoneRecognizer())
        engine.registry.register(BPJSRecognizerLegacy())
        engine.registry.register(DateOfBirthRecognizer())
        engine.registry.register(IndonesianPersonRecognizer())
        engine.initialize()
        self._engine = engine

        from sandiraksa.detection.context import DetectionContext, DetectionConfig

        def _ctx(entities: set[str], min_conf: float = 0.5) -> DetectionContext:
            return DetectionContext(
                operation_id=str(uuid4()),
                file_id=str(uuid4()),
                project_id=self._project_id,
                config=DetectionConfig(
                    enabled_entity_types=set(entities),
                    min_confidence=min_conf,
                ),
            )

        self._ctx_structured = _ctx(_STRUCTURED_ENTITIES)
        self._ctx_pii_only = _ctx(_PII_ONLY_ENTITIES)
        self._ctx_full = _ctx(_FULL_CONTENT_ENTITIES)

    # --------------------------------------------------------------- scanning

    def _scan_file(
        self, file_path: str, selected_columns: list | None
    ) -> list[dict]:
        path = Path(file_path)
        suffix = path.suffix.lower()

        if suffix in (".xlsx", ".xls") and selected_columns:
            return self._scan_xlsx_columns(path, selected_columns)

        content = self._read_content(path, suffix)
        if not content or not content.strip():
            return []

        results = self._engine.analyze_text(content, self._ctx_full)
        return [
            {
                "entity_type": r.entity_type,
                "text": r.text,
                "start": r.start,
                "end": r.end,
                "score": r.score,
            }
            for r in results
        ]

    def _read_content(self, path: Path, suffix: str) -> str:
        if suffix == ".csv":
            from sandiraksa.documents.csv_handler import CSVHandler

            return CSVHandler().read_file(path)
        if suffix in (".xlsx", ".xls"):
            return self._read_xlsx_content(path)
        if suffix == ".docx":
            return self._read_docx_content(path)
        return path.read_text(encoding="utf-8", errors="ignore")

    def _scan_xlsx_columns(self, path: Path, selected_columns: list) -> list[dict]:
        import openpyxl

        findings: list[dict] = []
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            from sandiraksa.detection.custom_patterns import find_custom_matches
        except Exception:  # pragma: no cover - optional feature
            find_custom_matches = None

        try:
            columns_by_sheet: dict = {}
            for col in selected_columns:
                columns_by_sheet.setdefault(col.sheet_name, []).append(col)

            cells_processed = 0
            for sheet_name, cols in columns_by_sheet.items():
                if self._cancelled:
                    break
                if sheet_name not in wb.sheetnames:
                    continue

                sheet = wb[sheet_name]
                col_indices = {c.column_index for c in cols}
                col_headers = {c.column_index: c.header for c in cols}
                col_types = {c.column_index: c.suggested_type for c in cols}
                col_modes = {
                    c.column_index: getattr(c, "protection_mode", "full_cell")
                    for c in cols
                }

                for row_num, row in enumerate(
                    sheet.iter_rows(min_row=2, values_only=True), start=2
                ):
                    if self._cancelled:
                        break
                    for col_idx in col_indices:
                        if col_idx - 1 >= len(row):
                            continue
                        cell_value = row[col_idx - 1]
                        if cell_value is None:
                            continue

                        value_str = str(cell_value)
                        header = col_headers.get(col_idx, "")
                        suggested_type = col_types.get(col_idx)
                        mode = col_modes.get(col_idx, "full_cell")
                        ctx_label = f"{sheet_name}!{header} (Row {row_num})"

                        if mode == "pii_only":
                            suggested_type = None

                        if suggested_type in _DIRECT_TYPES:
                            findings.append(
                                self._mk(
                                    suggested_type, value_str, 0, len(value_str),
                                    0.95, ctx_label, sheet_name, header, row_num,
                                )
                            )
                        else:
                            ctx = (
                                self._ctx_pii_only
                                if mode == "pii_only"
                                else self._ctx_structured
                            )
                            results = list(
                                self._engine.analyze_text(value_str, ctx)
                            )
                            if find_custom_matches is not None:
                                try:
                                    for cm in find_custom_matches(value_str):
                                        results.append(
                                            _Match(
                                                cm.label, cm.value,
                                                cm.start, cm.end, 0.95,
                                            )
                                        )
                                except Exception as ce:
                                    print(f"Custom pattern chain failed: {ce}")

                            if results:
                                for r in results:
                                    findings.append(
                                        self._mk(
                                            r.entity_type, r.text, r.start, r.end,
                                            r.score, ctx_label, sheet_name,
                                            header, row_num,
                                        )
                                    )
                            elif suggested_type:
                                findings.append(
                                    self._mk(
                                        suggested_type, value_str, 0,
                                        len(value_str), 0.8, ctx_label,
                                        sheet_name, header, row_num,
                                    )
                                )

                        cells_processed += 1
                        if cells_processed % _GC_EVERY_CELLS == 0:
                            gc.collect()
        finally:
            wb.close()

        return findings

    @staticmethod
    def _mk(
        entity_type, text, start, end, score,
        context, sheet, column, row,
    ) -> dict:
        return {
            "entity_type": entity_type,
            "text": text,
            "start": start,
            "end": end,
            "score": score,
            "context": context,
            "sheet": sheet,
            "column": column,
            "row": row,
        }

    def _read_xlsx_content(self, path: Path) -> str:
        try:
            import openpyxl

            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            parts: list[str] = []
            try:
                for sheet_name in wb.sheetnames:
                    sheet = wb[sheet_name]
                    for row in sheet.iter_rows(values_only=True):
                        for cell in row:
                            if cell is not None:
                                parts.append(str(cell))
            finally:
                wb.close()
            return "\n".join(parts)
        except ImportError:
            print("WARNING: openpyxl not installed, cannot read XLSX")
            return ""
        except Exception as exc:
            print(f"Error reading XLSX: {exc}")
            return ""

    def _read_docx_content(self, path: Path) -> str:
        try:
            from docx import Document

            doc = Document(path)
            parts: list[str] = []
            for para in doc.paragraphs:
                if para.text.strip():
                    parts.append(para.text)

            for table in doc.tables:
                rows = list(table.rows)
                if not rows:
                    continue
                headers = [c.text.strip() for c in rows[0].cells]
                for row in rows[1:]:
                    for col_idx, cell in enumerate(row.cells):
                        value = cell.text.strip()
                        if not value:
                            continue
                        header = headers[col_idx] if col_idx < len(headers) else ""
                        parts.append(f"{header}: {value}" if header else value)

            return "\n".join(parts)
        except ImportError:
            print("WARNING: python-docx not installed, cannot read DOCX")
            return ""
        except Exception as exc:
            print(f"Error reading DOCX: {exc}")
            return ""

    # ------------------------------------------------------------- finalize

    def _finalize_findings(
        self, findings: list[dict], file_id: str, file_name: str
    ) -> list[dict]:
        """Group by type, cap per type / total, and tag with ids/file info."""
        from collections import defaultdict

        by_type: dict[str, list[dict]] = defaultdict(list)
        for f in findings:
            by_type[f["entity_type"]].append(f)

        selected: list[dict] = []
        for _entity_type, type_findings in by_type.items():
            for finding in type_findings[:_MAX_PER_TYPE]:
                if len(selected) >= _MAX_TOTAL:
                    break
                finding["id"] = str(uuid4())
                finding["file_id"] = file_id
                finding["file_name"] = file_name
                finding["original"] = finding.get("text", "")
                finding["replacement"] = f"[{finding.get('entity_type', 'PII')}]"
                finding.setdefault("context", "")
                selected.append(finding)

        return selected


class _Match:
    """Lightweight match record for chained custom-pattern results."""

    __slots__ = ("entity_type", "text", "start", "end", "score")

    def __init__(self, entity_type, text, start, end, score) -> None:
        self.entity_type = entity_type
        self.text = text
        self.start = start
        self.end = end
        self.score = score


__all__ = ["ScanWorker"]
