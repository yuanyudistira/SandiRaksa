"""Tests for DOCX handler."""

import pytest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import MagicMock, patch

from sandiraksa.documents.docx_handler import (
    DOCXHandler,
    DOCXReader,
    DOCXWriter,
    DOCXMetadata,
    ParagraphInfo,
    TextRun,
    TextSpan,
    DocumentInventory,
    DOCXHandlerError,
    DocumentOpenError,
)


class TestTextRun:
    """Tests for TextRun dataclass."""

    def test_create(self):
        """Should create text run."""
        run = TextRun(
            text="Hello",
            run_index=0,
            start_offset=0,
            end_offset=5,
            is_bold=True,
        )
        assert run.text == "Hello"
        assert run.run_index == 0
        assert run.is_bold is True


class TestParagraphInfo:
    """Tests for ParagraphInfo dataclass."""

    def test_to_location(self):
        """Should convert to DocumentLocation."""
        para = ParagraphInfo(
            paragraph_index=5,
            text="Test paragraph",
            runs=[],
            component_type="paragraph",
        )
        loc = para.to_location("file-123")
        assert loc.file_id == "file-123"
        assert loc.component_type == "paragraph"
        assert loc.paragraph_index == 5

    def test_run_indices(self):
        """Should return list of run indices."""
        para = ParagraphInfo(
            paragraph_index=0,
            text="Hello World",
            runs=[
                TextRun("Hello ", 0, 0, 6),
                TextRun("World", 1, 6, 11),
            ],
        )
        assert para.run_indices == [0, 1]


class TestTextSpan:
    """Tests for TextSpan dataclass."""

    def test_to_location(self):
        """Should convert to DocumentLocation with run info."""
        span = TextSpan(
            text="test",
            paragraph_index=2,
            start_offset=10,
            end_offset=14,
            run_spans=[(0, 10, 14)],
        )
        loc = span.to_location()
        assert loc.paragraph_index == 2
        assert loc.start_offset == 10
        assert loc.end_offset == 14
        assert loc.run_indices == [0]


class TestDOCXReaderWithMock:
    """Tests for DOCXReader with mocked python-docx."""

    @pytest.fixture
    def mock_document(self):
        """Create mock document."""
        doc = MagicMock()
        
        # Mock paragraphs
        para1 = MagicMock()
        para1.text = "First paragraph"
        para1.runs = []
        para1.style.name = "Normal"
        
        para2 = MagicMock()
        para2.text = "Second paragraph"
        para2.runs = []
        para2.style.name = "Normal"
        
        doc.paragraphs = [para1, para2]
        
        # Mock tables
        doc.tables = []
        
        # Mock sections
        section = MagicMock()
        section.header.paragraphs = []
        section.footer.paragraphs = []
        doc.sections = [section]
        
        # Mock core properties
        props = MagicMock()
        props.author = "Test Author"
        props.title = "Test Document"
        props.created = None
        props.modified = None
        props.subject = None
        doc.core_properties = props
        
        return doc

    def test_paragraph_count(self, mock_document):
        """Should count paragraphs."""
        with patch("docx.Document", return_value=mock_document):
            reader = DOCXReader(Path("test.docx"))
            reader.open()
            
            paras = list(reader.iter_paragraphs())
            
            assert len(paras) == 2
            assert paras[0].text == "First paragraph"


class TestDOCXHandler:
    """Tests for high-level DOCXHandler."""

    def test_file_format(self):
        """Should return DOCX format."""
        from sandiraksa.domain.file_record import FileFormat
        
        handler = DOCXHandler(Path("test.docx"))
        assert handler.file_format == FileFormat.DOCX

    def test_default_output_path(self):
        """Should generate _protected suffix."""
        handler = DOCXHandler(Path("/path/to/document.docx"))
        output = handler._get_default_output_path()
        assert output == Path("/path/to/document_protected.docx")


class TestDOCXReaderMapOffsetToRuns:
    """Tests for offset to run mapping."""

    def test_single_run_mapping(self):
        """Should map offset within single run."""
        reader = DOCXReader(Path("test.docx"))
        
        para = ParagraphInfo(
            paragraph_index=0,
            text="Hello World",
            runs=[TextRun("Hello World", 0, 0, 11)],
        )
        
        # Map "World" (offset 6-11)
        spans = reader.map_offset_to_runs(para, 6, 11)
        
        assert len(spans) == 1
        assert spans[0] == (0, 6, 11)  # (run_idx, start_in_run, end_in_run)

    def test_multi_run_mapping(self):
        """Should map offset across multiple runs."""
        reader = DOCXReader(Path("test.docx"))
        
        para = ParagraphInfo(
            paragraph_index=0,
            text="Hello World!",
            runs=[
                TextRun("Hello ", 0, 0, 6),
                TextRun("World", 1, 6, 11),
                TextRun("!", 2, 11, 12),
            ],
        )
        
        # Map "lo Wor" (offset 3-9) - crosses two runs
        spans = reader.map_offset_to_runs(para, 3, 9)
        
        assert len(spans) == 2
        assert spans[0] == (0, 3, 6)  # "lo " in first run
        assert spans[1] == (1, 0, 3)  # "Wor" in second run

    def test_exact_run_boundary(self):
        """Should handle exact run boundaries."""
        reader = DOCXReader(Path("test.docx"))
        
        para = ParagraphInfo(
            paragraph_index=0,
            text="AB",
            runs=[
                TextRun("A", 0, 0, 1),
                TextRun("B", 1, 1, 2),
            ],
        )
        
        # Map just "A" (offset 0-1)
        spans = reader.map_offset_to_runs(para, 0, 1)
        
        assert len(spans) == 1
        assert spans[0] == (0, 0, 1)


class TestIntegration:
    """Integration tests requiring actual DOCX files."""

    @pytest.fixture
    def sample_docx(self):
        """Create a sample DOCX file."""
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        with TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.docx"
            
            doc = Document()
            
            # Add paragraphs
            doc.add_paragraph("John Doe is an employee.")
            doc.add_paragraph("Email: john@example.com")
            doc.add_paragraph("Phone: +62812345678")
            
            # Add a table
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Name"
            table.cell(0, 1).text = "Email"
            table.cell(1, 0).text = "Jane Smith"
            table.cell(1, 1).text = "jane@example.com"
            
            doc.save(path)
            
            yield path

    def test_read_paragraphs(self, sample_docx):
        """Should read paragraphs from file."""
        handler = DOCXHandler(sample_docx)
        
        paras = list(handler.iter_paragraphs(include_tables=False))
        
        assert len(paras) >= 3
        assert "John Doe" in paras[0].text
        assert "john@example.com" in paras[1].text

    def test_read_with_tables(self, sample_docx):
        """Should include table cell paragraphs."""
        handler = DOCXHandler(sample_docx)
        
        paras = list(handler.iter_paragraphs(include_tables=True))
        
        # Should include body paragraphs + table cells
        texts = [p.text for p in paras]
        assert any("Jane Smith" in t for t in texts)
        assert any("jane@example.com" in t for t in texts)

    def test_get_inventory(self, sample_docx):
        """Should get document inventory."""
        handler = DOCXHandler(sample_docx)
        
        inventory = handler.get_inventory()
        
        assert inventory.paragraph_count >= 3
        assert inventory.table_count == 1

    def test_get_metadata(self, sample_docx):
        """Should get document metadata."""
        handler = DOCXHandler(sample_docx)
        
        metadata = handler.get_metadata()
        
        assert metadata.paragraph_count >= 3
        assert metadata.table_count == 1

    def test_iter_text_segments(self, sample_docx):
        """Should iterate text segments."""
        handler = DOCXHandler(sample_docx)
        
        segments = list(handler.iter_text_segments())
        
        # Should have segments for each paragraph with text
        assert len(segments) >= 3
        
        # Check segment structure
        assert segments[0].text
        assert segments[0].location is not None

    def test_apply_treatments(self, sample_docx):
        """Should apply treatments and save."""
        with TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "output.docx"
            
            handler = DOCXHandler(sample_docx)
            
            treatments = [
                {
                    "paragraph_index": 0,
                    "old_text": "John Doe",
                    "new_text": "[[PERSON_ABC123]]",
                    "component_type": "paragraph",
                },
            ]
            
            result_path, modified = handler.apply_treatments(
                treatments,
                output_path=output_path,
            )
            
            assert result_path == output_path
            assert modified == 1
            assert output_path.exists()
            
            # Verify modifications
            from docx import Document
            doc = Document(output_path)
            assert "[[PERSON_ABC123]]" in doc.paragraphs[0].text

    def test_replace_across_runs(self, sample_docx):
        """Should replace text across multiple runs."""
        from docx import Document
        
        with TemporaryDirectory() as tmpdir:
            # Create document with multiple runs
            multi_run_path = Path(tmpdir) / "multi_run.docx"
            doc = Document()
            para = doc.add_paragraph()
            # Add text in multiple runs with different formatting
            run1 = para.add_run("John ")
            run1.bold = True
            run2 = para.add_run("Doe")
            run2.italic = True
            doc.save(multi_run_path)
            
            # Apply treatment
            output_path = Path(tmpdir) / "output.docx"
            handler = DOCXHandler(multi_run_path)
            
            treatments = [
                {
                    "paragraph_index": 0,
                    "old_text": "John Doe",
                    "new_text": "[[PERSON_XYZ]]",
                },
            ]
            
            result_path, modified = handler.apply_treatments(
                treatments,
                output_path=output_path,
            )
            
            assert modified == 1
            
            # Verify
            result_doc = Document(output_path)
            assert "[[PERSON_XYZ]]" in result_doc.paragraphs[0].text
