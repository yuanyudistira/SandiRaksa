"""Tests for XLSX handler."""

import pytest
from pathlib import Path
from tempfile import NamedTemporaryFile, TemporaryDirectory
from unittest.mock import MagicMock, patch

from sandiraksa.documents.xlsx_handler import (
    XLSXHandler,
    XLSXReader,
    XLSXWriter,
    XLSXCell,
    XLSXMetadata,
    SheetInventory,
    WorkbookInventory,
    XLSXHandlerError,
    WorkbookOpenError,
    _get_column_letter,
    _get_column_index,
)


class TestColumnConversion:
    """Tests for column letter/index conversion."""

    def test_column_letter_single(self):
        """Should convert single digit columns."""
        assert _get_column_letter(1) == "A"
        assert _get_column_letter(26) == "Z"

    def test_column_letter_double(self):
        """Should convert double digit columns."""
        assert _get_column_letter(27) == "AA"
        assert _get_column_letter(28) == "AB"
        assert _get_column_letter(52) == "AZ"
        assert _get_column_letter(53) == "BA"

    def test_column_letter_triple(self):
        """Should convert triple digit columns."""
        assert _get_column_letter(703) == "AAA"

    def test_column_index_single(self):
        """Should convert single letter columns."""
        assert _get_column_index("A") == 1
        assert _get_column_index("Z") == 26

    def test_column_index_double(self):
        """Should convert double letter columns."""
        assert _get_column_index("AA") == 27
        assert _get_column_index("AB") == 28
        assert _get_column_index("AZ") == 52
        assert _get_column_index("BA") == 53

    def test_column_index_case_insensitive(self):
        """Should handle lowercase."""
        assert _get_column_index("a") == 1
        assert _get_column_index("aa") == 27


class TestXLSXCell:
    """Tests for XLSXCell dataclass."""

    def test_coordinate(self):
        """Should return Excel coordinate."""
        cell = XLSXCell(
            sheet_name="Sheet1",
            row=5,
            column=3,
            column_letter="C",
            value="test",
            raw_value="test",
            cell_type="string",
        )
        assert cell.coordinate == "C5"

    def test_to_location(self):
        """Should convert to DocumentLocation."""
        cell = XLSXCell(
            sheet_name="Data",
            row=10,
            column=2,
            column_letter="B",
            value="hello",
            raw_value="hello",
            cell_type="string",
        )
        loc = cell.to_location()
        assert loc.component_type == "cell"
        assert loc.sheet_name == "Data"
        assert loc.cell_address == "B10"


class TestXLSXReaderWithMock:
    """Tests for XLSXReader with mocked openpyxl."""

    @pytest.fixture
    def mock_workbook(self):
        """Create mock workbook."""
        wb = MagicMock()
        wb.sheetnames = ["Sheet1", "Sheet2"]
        
        # Mock properties
        props = MagicMock()
        props.created = "2024-01-01"
        props.modified = "2024-01-02"
        props.creator = "Test User"
        props.title = "Test Workbook"
        wb.properties = props
        
        return wb

    @pytest.fixture
    def mock_worksheet(self):
        """Create mock worksheet."""
        ws = MagicMock()
        ws.min_row = 1
        ws.max_row = 10
        ws.min_column = 1
        ws.max_column = 5
        
        # Mock cells
        cell1 = MagicMock()
        cell1.value = "Header"
        cell1.row = 1
        cell1.column = 1
        cell1.data_type = "s"
        cell1.comment = None
        cell1.hyperlink = None
        
        cell2 = MagicMock()
        cell2.value = 123
        cell2.row = 2
        cell2.column = 1
        cell2.data_type = "n"
        cell2.comment = None
        cell2.hyperlink = None
        
        ws.iter_rows.return_value = [[cell1], [cell2]]
        
        return ws

    def test_sheet_names(self, mock_workbook):
        """Should return sheet names."""
        with patch("openpyxl.load_workbook", return_value=mock_workbook):
            with XLSXReader(Path("test.xlsx")) as reader:
                assert reader.sheet_names == ["Sheet1", "Sheet2"]

    def test_context_manager(self, mock_workbook):
        """Should work as context manager."""
        with patch("openpyxl.load_workbook", return_value=mock_workbook):
            with XLSXReader(Path("test.xlsx")) as reader:
                assert reader.workbook is not None
            # Should close
            mock_workbook.close.assert_called()


class TestXLSXWriterWithMock:
    """Tests for XLSXWriter with mocked openpyxl."""

    @pytest.fixture
    def mock_workbook(self):
        """Create mock workbook for writing."""
        wb = MagicMock()
        wb.sheetnames = ["Sheet1"]
        
        ws = MagicMock()
        cell = MagicMock()
        cell.data_type = "s"
        ws.cell.return_value = cell
        wb.__getitem__ = MagicMock(return_value=ws)
        
        return wb

    def test_set_cell_value(self, mock_workbook):
        """Should set cell value."""
        with patch("openpyxl.load_workbook", return_value=mock_workbook):
            writer = XLSXWriter(Path("test.xlsx"))
            writer.open()
            
            writer.set_cell_value("Sheet1", 1, 1, "new value")
            
            assert "Sheet1!A1" in writer._modified_cells

    def test_set_cell_by_coordinate(self, mock_workbook):
        """Should set cell by Excel coordinate."""
        with patch("openpyxl.load_workbook", return_value=mock_workbook):
            writer = XLSXWriter(Path("test.xlsx"))
            writer.open()
            
            writer.set_cell_by_coordinate("Sheet1", "B5", "test")
            
            assert "Sheet1!B5" in writer._modified_cells

    def test_apply_treatments(self, mock_workbook):
        """Should apply multiple treatments."""
        with patch("openpyxl.load_workbook", return_value=mock_workbook):
            writer = XLSXWriter(Path("test.xlsx"))
            writer.open()
            
            treatments = [
                {"sheet_name": "Sheet1", "row": 1, "column": 1, "new_value": "a"},
                {"sheet_name": "Sheet1", "row": 2, "column": 1, "new_value": "b"},
            ]
            
            modified = writer.apply_treatments(treatments)
            
            assert modified == 2


class TestXLSXHandler:
    """Tests for high-level XLSXHandler."""

    def test_file_format(self):
        """Should return XLSX format."""
        from sandiraksa.domain.file_record import FileFormat
        
        handler = XLSXHandler(Path("test.xlsx"))
        assert handler.file_format == FileFormat.XLSX

    def test_default_output_path(self):
        """Should generate _protected suffix."""
        handler = XLSXHandler(Path("/path/to/document.xlsx"))
        output = handler._get_default_output_path()
        assert output == Path("/path/to/document_protected.xlsx")


class TestIntegration:
    """Integration tests requiring actual Excel files."""

    @pytest.fixture
    def sample_xlsx(self):
        """Create a sample XLSX file."""
        try:
            from openpyxl import Workbook
        except ImportError:
            pytest.skip("openpyxl not installed")

        with TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "sample.xlsx"
            
            wb = Workbook()
            ws = wb.active
            ws.title = "TestSheet"
            
            # Add some data
            ws["A1"] = "Name"
            ws["B1"] = "Email"
            ws["C1"] = "Phone"
            
            ws["A2"] = "John Doe"
            ws["B2"] = "john@example.com"
            ws["C2"] = "+62812345678"
            
            ws["A3"] = "Jane Smith"
            ws["B3"] = "jane@example.com"
            ws["C3"] = "+62887654321"
            
            wb.save(path)
            wb.close()
            
            yield path

    def test_read_cells(self, sample_xlsx):
        """Should read cells from file."""
        handler = XLSXHandler(sample_xlsx)
        
        cells = list(handler.iter_cells())
        
        assert len(cells) == 9  # 3 rows x 3 columns
        assert cells[0].value == "Name"
        assert cells[1].value == "Email"
        assert cells[3].value == "John Doe"

    def test_get_inventory(self, sample_xlsx):
        """Should get workbook inventory."""
        handler = XLSXHandler(sample_xlsx)
        
        inventory = handler.get_inventory()
        
        assert len(inventory.sheets) == 1
        assert inventory.sheets[0].name == "TestSheet"
        assert inventory.sheets[0].row_count >= 3

    def test_get_metadata(self, sample_xlsx):
        """Should get workbook metadata."""
        handler = XLSXHandler(sample_xlsx)
        
        metadata = handler.get_metadata()
        
        assert metadata.sheet_count == 1
        assert "TestSheet" in metadata.sheet_names

    def test_iter_text_segments(self, sample_xlsx):
        """Should iterate text segments."""
        handler = XLSXHandler(sample_xlsx)
        
        segments = list(handler.iter_text_segments())
        
        # Should have segments for each cell with text
        assert len(segments) >= 9
        
        # Check segment structure
        segment = segments[0]
        assert segment.text == "Name"
        assert segment.location is not None
        assert segment.location.sheet_name == "TestSheet"

    def test_apply_treatments(self, sample_xlsx):
        """Should apply treatments and save."""
        with TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "output.xlsx"
            
            handler = XLSXHandler(sample_xlsx)
            
            treatments = [
                {"sheet_name": "TestSheet", "row": 2, "column": 1, "new_value": "[[PERSON_ABC123]]"},
                {"sheet_name": "TestSheet", "row": 2, "column": 2, "new_value": "[[EMAIL_DEF456]]"},
            ]
            
            result_path, modified = handler.apply_treatments(
                treatments,
                output_path=output_path,
            )
            
            assert result_path == output_path
            assert modified == 2
            assert output_path.exists()
            
            # Verify modifications
            verify_handler = XLSXHandler(output_path)
            cells = {
                f"{c.coordinate}": c.value 
                for c in verify_handler.iter_cells()
            }
            assert cells.get("A2") == "[[PERSON_ABC123]]"
            assert cells.get("B2") == "[[EMAIL_DEF456]]"
